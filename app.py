from flask import Flask, render_template, request, redirect, url_for, flash, send_file
import sqlite3, os, io
from datetime import datetime
from openpyxl import load_workbook, Workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'monitoring.db')
app = Flask(__name__)
app.secret_key = os.environ.get('APP_SECRET', 'change-this-secret-before-deployment')

SCHEMA = '''
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS districts (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 name TEXT NOT NULL UNIQUE,
 notes TEXT DEFAULT ''
);
CREATE TABLE IF NOT EXISTS schools (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 name TEXT NOT NULL,
 district_id INTEGER NOT NULL,
 address TEXT DEFAULT '',
 active INTEGER NOT NULL DEFAULT 1,
 UNIQUE(name, district_id),
 FOREIGN KEY(district_id) REFERENCES districts(id)
);
CREATE TABLE IF NOT EXISTS teachers (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 teacher_code TEXT UNIQUE,
 full_name TEXT NOT NULL,
 mother_name TEXT DEFAULT '',
 birth_date TEXT DEFAULT '',
 phone TEXT DEFAULT '',
 district_id INTEGER,
 school_id INTEGER,
 stage TEXT DEFAULT '',
 specialty TEXT DEFAULT '',
 gender TEXT DEFAULT '',
 employment_status TEXT DEFAULT 'على رأس العمل',
 notes TEXT DEFAULT '',
 active INTEGER NOT NULL DEFAULT 1,
 created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
 FOREIGN KEY(district_id) REFERENCES districts(id),
 FOREIGN KEY(school_id) REFERENCES schools(id)
);
CREATE TABLE IF NOT EXISTS import_log (
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 filename TEXT NOT NULL,
 imported_count INTEGER NOT NULL DEFAULT 0,
 skipped_count INTEGER NOT NULL DEFAULT 0,
 created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
 notes TEXT DEFAULT ''
);
'''

def connect():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys = ON')
    return conn

def init_db():
    with connect() as conn:
        conn.executescript(SCHEMA)
        # Seed a small set of illustrative districts only if database is empty.
        count = conn.execute('SELECT COUNT(*) FROM districts').fetchone()[0]
        if count == 0:
            conn.executemany('INSERT INTO districts(name) VALUES (?)', [('منطقة تجريبية 1',), ('منطقة تجريبية 2',)])

@app.route('/')
def index():
    with connect() as conn:
        stats = {
            'teachers': conn.execute('SELECT COUNT(*) FROM teachers WHERE active=1').fetchone()[0],
            'schools': conn.execute('SELECT COUNT(*) FROM schools WHERE active=1').fetchone()[0],
            'districts': conn.execute('SELECT COUNT(*) FROM districts').fetchone()[0],
            'duplicates': conn.execute('''SELECT COUNT(*) FROM teachers t WHERE t.full_name<>'' AND EXISTS (
                SELECT 1 FROM teachers t2 WHERE t2.id<t.id AND lower(trim(t2.full_name))=lower(trim(t.full_name))
                AND COALESCE(t2.birth_date,'')=COALESCE(t.birth_date,'')
            )''').fetchone()[0]
        }
        recent = conn.execute('''SELECT t.*, d.name district_name, s.name school_name
            FROM teachers t LEFT JOIN districts d ON d.id=t.district_id
            LEFT JOIN schools s ON s.id=t.school_id ORDER BY t.id DESC LIMIT 12''').fetchall()
        districts = conn.execute('SELECT * FROM districts ORDER BY name').fetchall()
    return render_template('index.html', stats=stats, recent=recent, districts=districts)

@app.route('/teachers')
def teachers():
    q = request.args.get('q', '').strip()
    with connect() as conn:
        rows = conn.execute('''SELECT t.*, d.name district_name, s.name school_name
            FROM teachers t LEFT JOIN districts d ON d.id=t.district_id
            LEFT JOIN schools s ON s.id=t.school_id
            WHERE (?='' OR t.full_name LIKE ? OR COALESCE(t.teacher_code,'') LIKE ? OR COALESCE(s.name,'') LIKE ?)
            ORDER BY t.full_name LIMIT 500''', (q, f'%{q}%', f'%{q}%', f'%{q}%')).fetchall()
    return render_template('teachers.html', rows=rows, q=q)

@app.route('/districts', methods=['GET', 'POST'])
def districts():
    with connect() as conn:
        if request.method == 'POST':
            name = request.form.get('name', '').strip()
            notes = request.form.get('notes', '').strip()
            if not name:
                flash('يرجى إدخال اسم المنطقة.', 'error')
            else:
                try:
                    conn.execute('INSERT INTO districts(name, notes) VALUES (?, ?)', (name, notes))
                    conn.commit()
                    flash('تمت إضافة المنطقة.', 'success')
                except sqlite3.IntegrityError:
                    flash('اسم المنطقة موجود مسبقًا.', 'error')
            return redirect(url_for('districts'))
        rows = conn.execute('''SELECT d.*, COUNT(DISTINCT s.id) school_count,
            COUNT(DISTINCT t.id) teacher_count FROM districts d
            LEFT JOIN schools s ON s.district_id=d.id LEFT JOIN teachers t ON t.district_id=d.id
            GROUP BY d.id ORDER BY d.name''').fetchall()
    return render_template('districts.html', rows=rows)

@app.route('/schools', methods=['GET', 'POST'])
def schools():
    with connect() as conn:
        if request.method == 'POST':
            name = request.form.get('name', '').strip()
            district_id = request.form.get('district_id', type=int)
            address = request.form.get('address', '').strip()
            if not name or not district_id:
                flash('يرجى إدخال اسم المدرسة واختيار المنطقة.', 'error')
            else:
                try:
                    conn.execute('INSERT INTO schools(name, district_id, address) VALUES (?, ?, ?)', (name, district_id, address))
                    conn.commit()
                    flash('تمت إضافة المدرسة.', 'success')
                except sqlite3.IntegrityError:
                    flash('المدرسة مسجلة مسبقًا ضمن هذه المنطقة.', 'error')
            return redirect(url_for('schools'))
        districts = conn.execute('SELECT * FROM districts ORDER BY name').fetchall()
        rows = conn.execute('''SELECT s.*, d.name district_name, COUNT(t.id) teacher_count
            FROM schools s JOIN districts d ON d.id=s.district_id
            LEFT JOIN teachers t ON t.school_id=s.id GROUP BY s.id ORDER BY d.name, s.name''').fetchall()
    return render_template('schools.html', rows=rows, districts=districts)

@app.route('/import', methods=['GET', 'POST'])
def import_excel():
    if request.method == 'POST':
        file = request.files.get('file')
        if not file or not file.filename.lower().endswith(('.xlsx', '.xlsm')):
            flash('اختر ملف Excel بصيغة .xlsx.', 'error')
            return redirect(url_for('import_excel'))
        try:
            wb = load_workbook(file, data_only=True, read_only=True)
            ws = wb.active
            raw_headers = [str(c.value or '').strip() for c in ws[1]]
            aliases = {
                'teacher_code': ['رقم المعلم', 'الرقم الوظيفي', 'معرف المعلم', 'teacher_id', 'teacher_code'],
                'full_name': ['الاسم الكامل', 'اسم المعلم', 'الاسم', 'full_name', 'name'],
                'mother_name': ['اسم الأم', 'اسم الام', 'mother_name'],
                'birth_date': ['تاريخ الميلاد', 'birth_date', 'dob'],
                'phone': ['رقم الهاتف', 'الهاتف', 'الجوال', 'phone'],
                'district': ['المنطقة التربوية', 'المنطقة', 'district'],
                'school': ['المدرسة', 'اسم المدرسة', 'school'],
                'stage': ['المرحلة التعليمية', 'المرحلة', 'stage'],
                'specialty': ['الاختصاص', 'التخصص', 'specialty'],
                'gender': ['الجنس', 'gender'],
                'employment_status': ['الحالة الوظيفية', 'employment_status'],
                'notes': ['ملاحظات', 'notes']
            }
            normalized = {}
            for idx, h in enumerate(raw_headers):
                for key, names in aliases.items():
                    if h.lower() in [n.lower() for n in names]:
                        normalized[key] = idx
            if 'full_name' not in normalized:
                flash('لم أجد عمود الاسم الكامل. استخدم نموذج الأعمدة المرفق في صفحة الاستيراد.', 'error')
                return redirect(url_for('import_excel'))
            imported = skipped = 0
            with connect() as conn:
                for row in ws.iter_rows(min_row=2, values_only=True):
                    def val(key):
                        i = normalized.get(key)
                        v = row[i] if i is not None and i < len(row) else ''
                        return '' if v is None else str(v).strip()
                    full_name = val('full_name')
                    if not full_name:
                        skipped += 1
                        continue
                    district_name = val('district') or 'غير محددة'
                    school_name = val('school') or 'غير محددة'
                    d = conn.execute('SELECT id FROM districts WHERE name=?', (district_name,)).fetchone()
                    if not d:
                        cur = conn.execute('INSERT INTO districts(name) VALUES (?)', (district_name,))
                        district_id = cur.lastrowid
                    else:
                        district_id = d['id']
                    s = conn.execute('SELECT id FROM schools WHERE name=? AND district_id=?', (school_name, district_id)).fetchone()
                    if not s:
                        cur = conn.execute('INSERT INTO schools(name, district_id) VALUES (?, ?)', (school_name, district_id))
                        school_id = cur.lastrowid
                    else:
                        school_id = s['id']
                    code = val('teacher_code') or None
                    existing = None
                    if code:
                        existing = conn.execute('SELECT id FROM teachers WHERE teacher_code=?', (code,)).fetchone()
                    if existing:
                        conn.execute('''UPDATE teachers SET full_name=?, mother_name=?, birth_date=?, phone=?, district_id=?, school_id=?, stage=?, specialty=?, gender=?, employment_status=?, notes=? WHERE id=?''',
                            (full_name, val('mother_name'), val('birth_date'), val('phone'), district_id, school_id, val('stage'), val('specialty'), val('gender'), val('employment_status') or 'على رأس العمل', val('notes'), existing['id']))
                    else:
                        conn.execute('''INSERT INTO teachers(teacher_code, full_name, mother_name, birth_date, phone, district_id, school_id, stage, specialty, gender, employment_status, notes)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                            (code, full_name, val('mother_name'), val('birth_date'), val('phone'), district_id, school_id, val('stage'), val('specialty'), val('gender'), val('employment_status') or 'على رأس العمل', val('notes')))
                    imported += 1
                conn.execute('INSERT INTO import_log(filename, imported_count, skipped_count) VALUES (?, ?, ?)', (file.filename, imported, skipped))
                conn.commit()
            flash(f'انتهى الاستيراد: {imported} سجلًا أُضيف أو حُدّث، و{skipped} صفًا تُرك لعدم وجود اسم.', 'success')
        except Exception as e:
            flash('تعذر قراءة الملف. تأكد من أنه ملف Excel صالح وأن الصف الأول يحتوي على أسماء الأعمدة.', 'error')
        return redirect(url_for('import_excel'))
    return render_template('import.html')

@app.route('/template')
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'بيانات المعلمين'
    headers = ['رقم المعلم', 'الاسم الكامل', 'اسم الأم', 'تاريخ الميلاد', 'رقم الهاتف', 'المنطقة التربوية', 'المدرسة', 'المرحلة التعليمية', 'الاختصاص', 'الجنس', 'الحالة الوظيفية', 'ملاحظات']
    ws.append(headers)
    ws.append(['T-0001', 'اسم المعلم الثلاثي أو الكامل', 'اسم الأم', '1990-01-01', '09xxxxxxxx', 'اسم المنطقة', 'اسم المدرسة', 'تعليم أساسي', 'رياضيات', 'ذكر', 'على رأس العمل', ''])
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = 22
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='نموذج_بيانات_المعلمين.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/health')
def health():
    return {'status': 'ok', 'app': 'Teacher Invigilation Management', 'copyright': 'jalal abboud 2026'}

if __name__ == '__main__':
    init_db()
    print('افتح المتصفح على: http://127.0.0.1:5000')
    app.run(host='0.0.0.0', port=5000, debug=False)
